"""Shared pytest fixtures: sample PDFs (latin + Chinese), docx, png."""
from __future__ import annotations

import os
import shutil
import tempfile
import time
import uuid
from pathlib import Path

import fitz
import pytest
from docx import Document
from PIL import Image

# Keep application-created uploads/sessions separate from the developer's
# backend/temp directory and remove this run's data when pytest exits.
_PYTEST_TEMP_ROOT = (
    Path(tempfile.gettempdir())
    / f"pdfed-{os.getpid():x}-{uuid.uuid4().hex[:6]}"
    / "app"
)
_PYTEST_TEMP_ROOT.mkdir(parents=True, exist_ok=True)
os.environ["PDF_EDITOR_TEMP_DIR"] = str(_PYTEST_TEMP_ROOT)
os.environ["PDF_EDITOR_MAX_UPLOAD_MB"] = "600"
os.environ["PDF_EDITOR_CORS_ORIGINS"] = "http://127.0.0.1:5173"


def make_sample_pdf(path: Path, pages: int = 3) -> Path:
    """3-page PDF, each page carrying 'PAGE-{n} HELLO WORLD'."""
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 100), f"PAGE-{i + 1} HELLO WORLD", fontsize=14)
    doc.save(str(path))
    doc.close()
    return path


def make_zh_pdf(path: Path) -> Path:
    """1-page PDF with Chinese text (via built-in china-s font)."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "你好世界 HELLO", fontsize=14, fontname="china-s")
    doc.save(str(path))
    doc.close()
    return path


def make_sample_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("测试文档", 0)
    doc.add_paragraph("这是用于 LibreOffice 转换测试的中文内容。")
    doc.save(str(path))
    return path


def make_sample_png(path: Path, size: tuple[int, int] = (120, 160)) -> Path:
    Image.new("RGB", size, (200, 60, 60)).save(str(path))
    return path


@pytest.fixture()
def sample_pdf(tmp_path: Path) -> Path:
    return make_sample_pdf(tmp_path / "sample.pdf")


@pytest.fixture()
def zh_pdf(tmp_path: Path) -> Path:
    return make_zh_pdf(tmp_path / "zh.pdf")


@pytest.fixture()
def sample_docx(tmp_path: Path) -> Path:
    return make_sample_docx(tmp_path / "sample.docx")


@pytest.fixture()
def sample_png(tmp_path: Path) -> Path:
    return make_sample_png(tmp_path / "sample.png")


@pytest.fixture(scope="session")
def client():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app, base_url="http://127.0.0.1") as c:
        yield c


@pytest.fixture()
def temp_dir_path():
    from app import config

    return config.TEMP_DIR


def submit_job(client, url: str, *, files=None, data=None, json=None, timeout=30):
    response = client.post(
        url,
        files=files,
        data=data,
        json=json,
        headers={"Idempotency-Key": uuid.uuid4().hex},
    )
    assert response.status_code == 202, response.text
    job_id = response.json()["jobId"]
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        status = client.get(f"/api/v1/jobs/{job_id}")
        assert status.status_code == 200, status.text
        result = status.json()
        if result["state"] in {"succeeded", "failed", "cancelled"}:
            file_response = (
                client.get(result["resultUrl"]) if result.get("resultUrl") else None
            )
            return result, file_response
        time.sleep(0.025)
    raise AssertionError(f"job {job_id} did not finish within {timeout}s")


def successful_job(client, url: str, **kwargs):
    result, file_response = submit_job(client, url, **kwargs)
    assert result["state"] == "succeeded", result.get("error")
    return result, file_response


@pytest.fixture(scope="session", autouse=True)
def cleanup_app_temp():
    yield
    try:
        from app.services.sessions import session_manager
        from app.core.tempfile import temp_manager

        session_manager.close_all()
        temp_manager.cleanup_old(0)
    finally:
        shutil.rmtree(_PYTEST_TEMP_ROOT.parent, ignore_errors=True)
